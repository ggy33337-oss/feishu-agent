# -*- coding: utf-8 -*-
from __future__ import annotations

import asyncio
import copy
import csv
import io
import re
import shutil
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.core.ports.storage import ArtifactStore
from config.settings import settings


class ExcelExecutor:
    """Structured XLSX/CSV runtime adapted from the MIT Hermes XLSX skill."""

    project_root = Path(__file__).resolve().parents[3]
    max_preview_rows = 50
    max_preview_columns = 30

    def __init__(self, artifact_store: ArtifactStore | None = None) -> None:
        self.artifact_store = artifact_store

    async def execute(self, inputs: dict[str, Any]) -> dict[str, Any]:
        return await asyncio.to_thread(self._execute_sync, inputs)

    def _execute_sync(self, inputs: dict[str, Any]) -> dict[str, Any]:
        parameters = inputs.get("parameters") if isinstance(inputs.get("parameters"), dict) else {}
        operation = str(parameters.get("operation") or "create").strip().lower()
        if operation in {"create", "csv_to_xlsx"}:
            parameters = self._normalize_create_parameters(parameters, str(inputs.get("source_text") or ""))
            operation = str(parameters.get("operation") or operation).strip().lower()
        task_id = self._safe_name(str(inputs.get("task_id") or "task"), "task")
        if self.artifact_store:
            output_dir = self.artifact_store.output_dir(task_id)
        else:
            artifact_root = Path(settings.artifact_root) if settings.artifact_root else self.project_root / "outputs"
            output_dir = artifact_root / task_id
            output_dir.mkdir(parents=True, exist_ok=True)

        if operation == "create":
            return self._create(
                self._preserve_source_table(parameters, str(inputs.get("source_text") or "")),
                output_dir,
            )
        if operation == "inspect":
            input_path = self._resolve_input_path(inputs, parameters)
            return {"operation": operation, "input_path": str(input_path), "inspection": self._inspect(input_path)}
        if operation == "edit":
            return self._edit(inputs, parameters, output_dir)
        if operation == "csv_to_xlsx":
            return self._csv_to_xlsx(inputs, parameters, output_dir)
        if operation == "xlsx_to_csv":
            return self._xlsx_to_csv(inputs, parameters, output_dir)
        raise ValueError(f"Unsupported Excel operation: {operation}")

    def _normalize_create_parameters(
        self,
        parameters: dict[str, Any],
        source_text: str,
    ) -> dict[str, Any]:
        """Prevent malformed model JSON from turning a valid create request into a dead job."""
        normalized = copy.deepcopy(parameters)
        source_rows = self._parse_source_table(source_text)
        operation = str(normalized.get("operation") or "create").strip().lower()
        if operation == "csv_to_xlsx" and source_rows:
            normalized["operation"] = "create"
            normalized["workbook_spec"] = {
                "sheets": [{"name": "数据", "rows": source_rows, "freeze_panes": "A2", "autofilter": True}]
            }
        filename = str(normalized.get("filename") or "示例表格.xlsx")
        normalized["filename"] = filename if filename.lower().endswith(".xlsx") else f"{filename}.xlsx"
        spec = normalized.get("workbook_spec")
        if not isinstance(spec, dict):
            spec = {}
        sheets = spec.get("sheets")
        if not isinstance(sheets, list) or not sheets:
            sheets = [{}]
        fallback_rows = source_rows or [
            ["项目", "数量", "单价", "金额"],
            ["A", 10, 25, ""],
            ["B", 5, 40, ""],
            ["C", 8, 15, ""],
            ["合计", "", "", ""],
        ]
        normalized_sheets: list[dict[str, Any]] = []
        for index, raw_sheet in enumerate(sheets, start=1):
            sheet = dict(raw_sheet) if isinstance(raw_sheet, dict) else {}
            rows = sheet.get("rows")
            if not isinstance(rows, list) or not all(isinstance(row, list) for row in rows):
                rows = []
            sheet["rows"] = rows or fallback_rows
            if not isinstance(sheet.get("cells"), dict):
                sheet["cells"] = {}
            sheet.setdefault("name", "数据" if index == 1 else f"数据{index}")
            sheet.setdefault("freeze_panes", "A2")
            sheet.setdefault("autofilter", True)
            normalized_sheets.append(sheet)
        spec["sheets"] = normalized_sheets
        normalized["workbook_spec"] = spec
        return normalized

    def _preserve_source_table(
        self,
        parameters: dict[str, Any],
        source_text: str,
    ) -> dict[str, Any]:
        source_rows = self._parse_source_table(source_text)
        if not source_rows:
            return parameters

        normalized = copy.deepcopy(parameters)
        workbook_spec = normalized.get("workbook_spec")
        if not isinstance(workbook_spec, dict):
            return normalized
        sheets = workbook_spec.get("sheets")
        if not isinstance(sheets, list) or not sheets or not isinstance(sheets[0], dict):
            return normalized

        sheet = sheets[0]
        sheet["rows"] = source_rows
        self._remove_source_value_overrides(sheet, len(source_rows), len(source_rows[0]))
        return normalized

    def _parse_source_table(self, source_text: str) -> list[list[Any]]:
        if not source_text.strip() or max(source_text.count(","), source_text.count("\t")) < 1:
            return []
        delimiter = "\t" if source_text.count("\t") > source_text.count(",") else ","
        parsed = [
            [value.strip() for value in row]
            for row in csv.reader(io.StringIO(source_text), delimiter=delimiter)
            if any(value.strip() for value in row)
        ]
        widths = Counter(len(row) for row in parsed if len(row) >= 2)
        if not widths:
            return []
        expected_width = max(widths, key=lambda width: (widths[width], width))
        rows = [row for row in parsed if len(row) == expected_width]
        if len(rows) < 2:
            return []

        headers = rows[0]
        return [headers] + [
            [self._typed_source_value(headers[index], value) for index, value in enumerate(row)]
            for row in rows[1:]
        ]

    def _typed_source_value(self, header: str, value: str) -> Any:
        if value == "":
            return None
        normalized_header = header.strip().lower()
        if any(keyword in normalized_header for keyword in ("日期", "date")):
            try:
                return date.fromisoformat(value)
            except ValueError:
                return value
        numeric_keywords = (
            "工资",
            "金额",
            "数量",
            "年龄",
            "价格",
            "成本",
            "收入",
            "利润",
            "salary",
            "amount",
            "quantity",
            "age",
            "price",
            "cost",
            "revenue",
            "profit",
        )
        if any(keyword in normalized_header for keyword in numeric_keywords):
            for caster in (int, float):
                try:
                    return caster(value)
                except ValueError:
                    pass
        return value

    def _remove_source_value_overrides(
        self,
        sheet: dict[str, Any],
        row_count: int,
        column_count: int,
    ) -> None:
        cells = sheet.get("cells")
        if not isinstance(cells, dict):
            return
        for coordinate in list(cells):
            try:
                min_column, min_row, max_column, max_row = range_boundaries(str(coordinate))
            except ValueError:
                continue
            overlaps_source = (
                min_row <= row_count
                and max_row >= 1
                and min_column <= column_count
                and max_column >= 1
            )
            if not overlaps_source:
                continue
            spec = cells[coordinate]
            if not isinstance(spec, dict):
                del cells[coordinate]
                continue
            spec.pop("value", None)
            spec.pop("formula", None)

    def _create(self, parameters: dict[str, Any], output_dir: Path) -> dict[str, Any]:
        spec = parameters.get("workbook_spec")
        if not isinstance(spec, dict) or not isinstance(spec.get("sheets"), list) or not spec["sheets"]:
            raise ValueError("Excel create operation requires workbook_spec.sheets.")
        output_path = output_dir / self._xlsx_name(parameters.get("filename"), "workbook.xlsx")
        workbook = Workbook()
        workbook.remove(workbook.active)
        for index, sheet_spec in enumerate(spec["sheets"], start=1):
            if not isinstance(sheet_spec, dict):
                raise ValueError(f"workbook_spec.sheets[{index - 1}] must be an object.")
            worksheet = workbook.create_sheet(self._sheet_name(sheet_spec.get("name"), f"Sheet{index}"))
            self._build_sheet(worksheet, sheet_spec)
        workbook.calculation.fullCalcOnLoad = bool(spec.get("full_calc_on_load", True))
        workbook.calculation.forceFullCalc = True
        workbook.save(output_path)
        return self._output_result("create", output_path)

    def _edit(self, inputs: dict[str, Any], parameters: dict[str, Any], output_dir: Path) -> dict[str, Any]:
        input_path = self._resolve_input_path(inputs, parameters)
        if input_path.suffix.lower() != ".xlsx":
            raise ValueError("Edit currently supports .xlsx files only.")
        output_path = output_dir / self._xlsx_name(parameters.get("filename"), f"edited-{input_path.name}")
        shutil.copy2(input_path, output_path)
        workbook = load_workbook(output_path)
        default_sheet = str(parameters.get("sheet_name") or workbook.active.title)
        if default_sheet not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {default_sheet}")
        sets = parameters.get("sets", {})
        if not isinstance(sets, dict):
            raise ValueError("Excel edit sets must map cell references to values.")
        for reference, value in sets.items():
            sheet_name, coordinate = self._split_reference(str(reference), default_sheet)
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet not found: {sheet_name}")
            self._apply_cell(workbook[sheet_name], coordinate, value)
        append_rows = parameters.get("append_rows", [])
        if not isinstance(append_rows, list):
            raise ValueError("Excel edit append_rows must be an array.")
        for row in append_rows:
            if not isinstance(row, list):
                raise ValueError("Each appended row must be an array.")
            workbook[default_sheet].append(row)
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(output_path)
        return self._output_result("edit", output_path)

    def _csv_to_xlsx(self, inputs: dict[str, Any], parameters: dict[str, Any], output_dir: Path) -> dict[str, Any]:
        input_path = self._resolve_input_path(inputs, parameters)
        if input_path.suffix.lower() not in {".csv", ".tsv"}:
            raise ValueError("csv_to_xlsx requires a .csv or .tsv input file.")
        delimiter = "\t" if input_path.suffix.lower() == ".tsv" else str(parameters.get("delimiter") or ",")
        with open(input_path, newline="", encoding="utf-8") as file_handle:
            rows = list(csv.reader(file_handle, delimiter=delimiter))
        output_path = output_dir / self._xlsx_name(parameters.get("filename"), f"{input_path.stem}.xlsx")
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = self._sheet_name(parameters.get("sheet_name"), "Data")
        for row_index, row in enumerate(rows):
            worksheet.append(row if row_index == 0 else [self._infer_csv_value(value) for value in row])
        self._apply_default_layout(worksheet)
        workbook.save(output_path)
        return self._output_result("csv_to_xlsx", output_path)

    def _xlsx_to_csv(self, inputs: dict[str, Any], parameters: dict[str, Any], output_dir: Path) -> dict[str, Any]:
        input_path = self._resolve_input_path(inputs, parameters)
        workbook = load_workbook(input_path, data_only=bool(parameters.get("data_only", True)))
        sheet_name = str(parameters.get("sheet_name") or workbook.active.title)
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {sheet_name}")
        output_name = self._safe_name(str(parameters.get("filename") or f"{input_path.stem}.csv"), "output.csv")
        if not output_name.lower().endswith(".csv"):
            output_name += ".csv"
        output_path = output_dir / output_name
        with open(output_path, "w", newline="", encoding="utf-8") as file_handle:
            writer = csv.writer(file_handle)
            for row in workbook[sheet_name].iter_rows(values_only=True):
                writer.writerow([self._jsonable(value) if value is not None else "" for value in row])
        return {
            "operation": "xlsx_to_csv",
            "output_path": str(output_path),
            "output_name": output_name,
            "mime_type": "text/csv",
            "verification": {"exists": output_path.exists(), "size": output_path.stat().st_size},
        }

    def _build_sheet(self, worksheet: Any, spec: dict[str, Any]) -> None:
        rows = spec.get("rows", [])
        if not isinstance(rows, list):
            raise ValueError("Worksheet rows must be an array.")
        for row in rows:
            if not isinstance(row, list):
                raise ValueError("Each worksheet row must be an array.")
            worksheet.append([None if isinstance(item, dict) else item for item in row])
            for column_number, item in enumerate(row, start=1):
                if isinstance(item, dict):
                    coordinate = worksheet.cell(row=worksheet.max_row, column=column_number).coordinate
                    self._apply_cell(worksheet, coordinate, item)
        cells = spec.get("cells", {})
        if not isinstance(cells, dict):
            raise ValueError("Worksheet cells must be an object.")
        for coordinate, cell_spec in cells.items():
            self._apply_cell(worksheet, str(coordinate), cell_spec)
        for column, width in self._mapping(spec.get("column_widths")).items():
            worksheet.column_dimensions[str(column)].width = min(float(width), 60)
        for row_number, height in self._mapping(spec.get("row_heights")).items():
            worksheet.row_dimensions[int(row_number)].height = min(float(height), 100)
        for cell_range in self._list(spec.get("merges")):
            worksheet.merge_cells(str(cell_range))
        worksheet.freeze_panes = spec.get("freeze_panes")
        autofilter = spec.get("autofilter")
        if autofilter is True:
            worksheet.auto_filter.ref = worksheet.dimensions
        elif isinstance(autofilter, str) and autofilter.strip():
            worksheet.auto_filter.ref = autofilter.strip()
        for item in self._list(spec.get("conditional_formats")):
            self._add_conditional_format(worksheet, item)
        for item in self._list(spec.get("validations")):
            self._add_validation(worksheet, item)
        for item in self._list(spec.get("tables")):
            self._add_table(worksheet, item)
        for item in self._list(spec.get("charts")):
            self._add_chart(worksheet, item)
        self._apply_default_layout(worksheet)

    def _apply_cell(self, worksheet: Any, coordinate: str, spec: Any) -> None:
        if ":" in coordinate:
            if not isinstance(spec, dict) or "value" in spec or "formula" in spec:
                raise ValueError(
                    f"Cell range {coordinate} may contain formatting only; values and formulas require single-cell references."
                )
            min_column, min_row, max_column, max_row = range_boundaries(coordinate)
            for row in worksheet.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_column,
                max_col=max_column,
            ):
                for cell in row:
                    self._apply_cell_style(cell, spec)
            return

        cell = worksheet[coordinate]
        if not isinstance(spec, dict):
            cell.value = spec
            return
        if "formula" in spec:
            formula = str(spec["formula"])
            cell.value = formula if formula.startswith("=") else f"={formula}"
        elif "value" in spec:
            cell.value = self._typed_value(spec["value"], spec.get("type"))
        if spec.get("hyperlink"):
            cell.hyperlink = str(spec["hyperlink"])
            cell.style = "Hyperlink"
        if spec.get("note"):
            note = spec["note"]
            if isinstance(note, dict):
                cell.comment = Comment(str(note.get("text") or ""), str(note.get("author") or "User"))
            else:
                cell.comment = Comment(str(note), "User")
        self._apply_cell_style(cell, spec)

    def _apply_cell_style(self, cell: Any, spec: dict[str, Any]) -> None:
        if spec.get("format"):
            cell.number_format = str(spec["format"])
        cell.font = Font(
            name=str(spec.get("font_name") or "Arial"),
            size=float(spec.get("font_size") or 10),
            bold=bool(spec.get("bold")),
            italic=bool(spec.get("italic")),
            color=spec.get("font_color"),
        )
        if spec.get("fill"):
            cell.fill = PatternFill("solid", fgColor=str(spec["fill"]))
        if spec.get("border"):
            side = Side(style=str(spec["border"]), color="D9D9D9")
            cell.border = Border(left=side, right=side, top=side, bottom=side)
        cell.alignment = Alignment(
            horizontal=spec.get("align"),
            vertical=spec.get("valign"),
            wrap_text=bool(spec.get("wrap")),
        )

    def _apply_default_layout(self, worksheet: Any) -> None:
        if worksheet.max_row >= 1:
            for cell in worksheet[1]:
                if cell.value is not None and not cell.has_style:
                    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1F4E78")
                    cell.alignment = Alignment(vertical="center")
            if worksheet.freeze_panes is None and worksheet.max_row > 1:
                worksheet.freeze_panes = "A2"
            if worksheet.auto_filter.ref is None and worksheet.max_row > 1:
                worksheet.auto_filter.ref = worksheet.dimensions
        for column_number in range(1, worksheet.max_column + 1):
            letter = get_column_letter(column_number)
            if worksheet.column_dimensions[letter].width not in (None, 13.0):
                continue
            longest = max(
                (len(str(worksheet.cell(row=row, column=column_number).value or "")) for row in range(1, worksheet.max_row + 1)),
                default=8,
            )
            worksheet.column_dimensions[letter].width = min(max(longest + 2, 10), 40)

    def _add_conditional_format(self, worksheet: Any, spec: Any) -> None:
        if not isinstance(spec, dict) or not spec.get("range"):
            raise ValueError("Conditional format requires range.")
        if spec.get("type") == "color_scale":
            rule = ColorScaleRule(
                start_type="min",
                start_color=str(spec.get("start_color") or "FFF8696B"),
                end_type="max",
                end_color=str(spec.get("end_color") or "FF63BE7B"),
            )
        elif spec.get("type") == "data_bar":
            color = str(spec.get("color") or "4472C4").lstrip("#")
            rule = DataBarRule(start_type="min", end_type="max", color=color)
        else:
            fill = PatternFill("solid", fgColor=str(spec.get("fill") or "FFC7CE"))
            formula = spec.get("formula") if isinstance(spec.get("formula"), list) else ["0"]
            rule = CellIsRule(operator=str(spec.get("operator") or "greaterThan"), formula=formula, fill=fill)
        worksheet.conditional_formatting.add(str(spec["range"]), rule)

    def _add_validation(self, worksheet: Any, spec: Any) -> None:
        if not isinstance(spec, dict) or not spec.get("range") or not spec.get("formula1"):
            raise ValueError("Data validation requires range and formula1.")
        validation = DataValidation(
            type=str(spec.get("type") or "list"),
            formula1=str(spec["formula1"]),
            allow_blank=bool(spec.get("allow_blank", True)),
        )
        validation.add(str(spec["range"]))
        worksheet.add_data_validation(validation)

    def _add_table(self, worksheet: Any, spec: Any) -> None:
        if not isinstance(spec, dict) or not spec.get("name") or not (spec.get("range") or spec.get("ref")):
            raise ValueError("Table requires name and range.")
        table_range = str(spec.get("range") or spec["ref"])
        table = Table(displayName=self._safe_name(str(spec["name"]), "Table1"), ref=table_range)
        table.tableStyleInfo = TableStyleInfo(
            name=str(spec.get("style") or "TableStyleMedium2"),
            showRowStripes=bool(spec.get("row_stripes", True)),
            showColumnStripes=bool(spec.get("column_stripes", False)),
        )
        worksheet.add_table(table)

    def _add_chart(self, worksheet: Any, spec: Any) -> None:
        if not isinstance(spec, dict) or not spec.get("data"):
            raise ValueError("Chart requires data range.")
        chart_classes = {"bar": BarChart, "line": LineChart, "pie": PieChart}
        chart_type = str(spec.get("type") or "bar").lower()
        if chart_type not in chart_classes:
            raise ValueError(f"Unsupported chart type: {chart_type}")
        chart = chart_classes[chart_type]()
        chart.title = str(spec.get("title") or "")
        chart.add_data(self._reference(worksheet, str(spec["data"])), titles_from_data=bool(spec.get("titles_from_data", True)))
        if spec.get("categories"):
            chart.set_categories(self._reference(worksheet, str(spec["categories"])))
        worksheet.add_chart(chart, str(spec.get("anchor") or "H2"))

    def _output_result(self, operation: str, output_path: Path) -> dict[str, Any]:
        inspection = self._inspect(output_path)
        return {
            "operation": operation,
            "output_path": str(output_path),
            "output_name": output_path.name,
            "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "verification": {
                "exists": output_path.exists(),
                "size": output_path.stat().st_size,
                "formula_errors": inspection["formula_errors"],
                "sheets": inspection["sheets"],
            },
            "inspection": inspection,
        }

    def _inspect(self, path: Path) -> dict[str, Any]:
        workbook = load_workbook(path, data_only=False, read_only=False)
        sheets: list[dict[str, Any]] = []
        previews: dict[str, list[list[Any]]] = {}
        formulas: list[dict[str, str]] = []
        formula_errors: list[dict[str, str]] = []
        for worksheet in workbook.worksheets:
            sheets.append(
                {
                    "name": worksheet.title,
                    "dimensions": worksheet.dimensions,
                    "rows": worksheet.max_row,
                    "columns": worksheet.max_column,
                    "charts": len(getattr(worksheet, "_charts", [])),
                    "tables": {name: worksheet.tables[name].ref for name in worksheet.tables.keys()},
                }
            )
            preview: list[list[Any]] = []
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=min(worksheet.max_row, self.max_preview_rows),
                max_col=min(worksheet.max_column, self.max_preview_columns),
            ):
                preview.append([self._jsonable(cell.value) for cell in row])
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas.append({"sheet": worksheet.title, "cell": cell.coordinate, "formula": cell.value})
                    if isinstance(cell.value, str) and re.search(r"#(?:REF!|DIV/0!|VALUE!|NAME\?|N/A)", cell.value):
                        formula_errors.append({"sheet": worksheet.title, "cell": cell.coordinate, "value": cell.value})
            previews[worksheet.title] = preview
        return {"path": str(path), "sheets": sheets, "previews": previews, "formulas": formulas, "formula_errors": formula_errors}

    def _resolve_input_path(self, inputs: dict[str, Any], parameters: dict[str, Any]) -> Path:
        candidates: list[Any] = [parameters.get("input_path")]
        attachments = inputs.get("attachments")
        if isinstance(attachments, list):
            for attachment in attachments:
                if isinstance(attachment, dict):
                    candidates.extend([attachment.get("local_path"), attachment.get("path")])
        for candidate in candidates:
            if candidate:
                path = Path(str(candidate)).expanduser().resolve()
                if path.is_file():
                    return path
        raise ValueError("Excel operation requires a readable input_path or downloaded attachment.")

    def _xlsx_name(self, value: Any, fallback: str) -> str:
        name = self._safe_name(str(value or fallback), fallback)
        return name if name.lower().endswith(".xlsx") else f"{name}.xlsx"

    def _safe_name(self, value: str, fallback: str) -> str:
        cleaned = re.sub(r"[^0-9A-Za-z._\-\u4e00-\u9fff]+", "_", value).strip("._")
        return cleaned[:100] or fallback

    def _sheet_name(self, value: Any, fallback: str) -> str:
        cleaned = re.sub(r"[\[\]:*?/\\]", "_", str(value or fallback)).strip()
        return cleaned[:31] or fallback

    def _split_reference(self, reference: str, default_sheet: str) -> tuple[str, str]:
        if "!" not in reference:
            return default_sheet, reference
        sheet_name, coordinate = reference.rsplit("!", 1)
        return sheet_name.strip("'"), coordinate

    def _reference(self, worksheet: Any, cell_range: str) -> Reference:
        structured_match = re.search(r"\[([$]?[A-Za-z]{1,3}[$]?\d+:[$]?[A-Za-z]{1,3}[$]?\d+)\]", cell_range)
        if structured_match:
            cell_range = structured_match.group(1)
        min_column, min_row, max_column, max_row = range_boundaries(cell_range)
        return Reference(worksheet, min_col=min_column, min_row=min_row, max_col=max_column, max_row=max_row)

    def _typed_value(self, value: Any, type_hint: Any) -> Any:
        if type_hint == "date" and isinstance(value, str):
            return date.fromisoformat(value)
        if type_hint == "datetime" and isinstance(value, str):
            return datetime.fromisoformat(value)
        return value

    def _infer_csv_value(self, value: str) -> Any:
        if value == "":
            return None
        if value.lower() in {"true", "false"}:
            return value.lower() == "true"
        for caster in (int, float):
            try:
                return caster(value)
            except ValueError:
                pass
        for parser in (date.fromisoformat, datetime.fromisoformat):
            try:
                return parser(value)
            except ValueError:
                pass
        return value

    def _jsonable(self, value: Any) -> Any:
        return value.isoformat() if isinstance(value, (datetime, date, time)) else value

    def _mapping(self, value: Any) -> dict[str, Any]:
        return value if isinstance(value, dict) else {}

    def _list(self, value: Any) -> list[Any]:
        return value if isinstance(value, list) else []
