# -*- coding: utf-8 -*-
import asyncio

import httpx
from openpyxl import load_workbook

from app.skills.registry import SkillRegistry
from app.skills.excel.executor import ExcelExecutor
from app.skills.excel.skill import ExcelSkill
from app.skills.search.executor import SearchExecutor
from config.settings import settings


def test_default_registry_contains_core_skills() -> None:
    registry = SkillRegistry.default()

    assert "excel" in registry.available_names()
    assert "feishu_lark_agent" not in registry.available_names()


def test_excel_skill_creates_verified_workbook(tmp_path) -> None:
    executor = ExcelExecutor()
    executor.project_root = tmp_path
    skill = ExcelSkill(executor)

    result = asyncio.run(
        skill.run(
            {
                "task_id": "excel-test",
                "parameters": {
                    "operation": "create",
                    "filename": "销售分析.xlsx",
                    "workbook_spec": {
                        "sheets": [
                            {
                                "name": "销售",
                                "rows": [
                                    ["产品", "收入", "成本", "利润"],
                                    ["A", 100, 60, {"formula": "B2-C2", "format": "#,##0"}],
                                    ["B", 200, 90, {"formula": "B3-C3", "format": "#,##0"}],
                                ],
                                "freeze_panes": "A2",
                                "autofilter": True,
                                "conditional_formats": [
                                    {"type": "data_bar", "range": "D2:D3", "color": "#4472C4"}
                                ],
                                "tables": [{"name": "SalesTable", "ref": "A1:D3"}],
                                "charts": [
                                    {
                                        "type": "bar",
                                        "title": "利润",
                                        "data": "SalesTable[D1:D3]",
                                        "categories": "SalesTable[A2:A3]",
                                        "anchor": "F2",
                                    }
                                ],
                            }
                        ]
                    },
                },
            }
        )
    )

    output_path = tmp_path / "outputs" / "excel-test" / "销售分析.xlsx"
    workbook = load_workbook(output_path, data_only=False)
    worksheet = workbook["销售"]
    assert result.ok is True
    assert result.artifacts[0]["path"] == str(output_path)
    assert result.data["verification"]["formula_errors"] == []
    assert worksheet["D2"].value == "=B2-C2"
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:D3"
    assert "SalesTable" in worksheet.tables
    assert len(worksheet._charts) == 1


def test_excel_skill_applies_formatting_to_cell_range(tmp_path) -> None:
    executor = ExcelExecutor()
    executor.project_root = tmp_path
    skill = ExcelSkill(executor)

    result = asyncio.run(
        skill.run(
            {
                "task_id": "excel-range-style",
                "parameters": {
                    "operation": "create",
                    "filename": "员工信息.xlsx",
                    "workbook_spec": {
                        "sheets": [
                            {
                                "name": "员工信息",
                                "rows": [
                                    ["工号", "姓名", "部门"],
                                    ["EMP001", "张伟", "技术部"],
                                ],
                                "cells": {
                                    "A1:C1": {
                                        "bold": True,
                                        "font_color": "FFFFFF",
                                        "fill": "1F4E78",
                                    }
                                },
                            }
                        ]
                    },
                },
            }
        )
    )

    workbook = load_workbook(result.artifacts[0]["path"])
    worksheet = workbook["员工信息"]
    assert all(worksheet.cell(1, column).font.bold for column in range(1, 4))
    assert all(worksheet.cell(1, column).fill.fgColor.rgb == "001F4E78" for column in range(1, 4))


def test_excel_skill_normalizes_malformed_model_rows(tmp_path) -> None:
    executor = ExcelExecutor()
    executor.project_root = tmp_path
    result = asyncio.run(
        ExcelSkill(executor).run(
            {
                "task_id": "excel-malformed-model",
                "parameters": {
                    "operation": "create",
                    "filename": "兜底.xlsx",
                    "workbook_spec": {"sheets": [{"name": "数据", "rows": 5}]},
                },
            }
        )
    )

    assert result.ok is True
    workbook = load_workbook(result.artifacts[0]["path"])
    assert workbook.active.max_row >= 2


def test_excel_skill_converts_inline_csv_when_model_selected_conversion(tmp_path) -> None:
    executor = ExcelExecutor()
    executor.project_root = tmp_path
    source_text = "name,department\n陈宇恒,研发部\n生成excel表"
    result = asyncio.run(
        ExcelSkill(executor).run(
            {
                "task_id": "excel-inline-csv",
                "source_text": source_text,
                "parameters": {"operation": "csv_to_xlsx"},
            }
        )
    )

    workbook = load_workbook(result.artifacts[0]["path"], data_only=False)
    assert workbook.active.cell(1, 1).value == "name"
    assert workbook.active.cell(2, 1).value == "陈宇恒"


def test_excel_skill_preserves_original_csv_rows_over_model_rewrite(tmp_path) -> None:
    executor = ExcelExecutor()
    executor.project_root = tmp_path
    skill = ExcelSkill(executor)
    source_text = (
        "工号,姓名,部门,基本工资(元),联系电话\n"
        "EMP001,张伟,技术部,18000,13800138001\n"
        "EMP002,李娜,产品部,16500,13800138002\n"
        "生成excel表"
    )

    result = asyncio.run(
        skill.run(
            {
                "task_id": "excel-source-preservation",
                "source_text": source_text,
                "parameters": {
                    "operation": "create",
                    "filename": "员工信息.xlsx",
                    "workbook_spec": {
                        "sheets": [
                            {
                                "name": "员工信息",
                                "rows": [
                                    ["Employee ID", "Name", "Department", "Salary", "Phone"],
                                    ["EMP001", "Zhang Wei", "Engineering", 18000, 13800138001],
                                    ["EMP002", "Li Na", "Product", 16500, 13800138002],
                                ],
                                "cells": {
                                    "A1:E1": {"bold": True, "fill": "1F4E78"},
                                    "B2": {"value": "Zhang Wei", "italic": True},
                                },
                            }
                        ]
                    },
                },
            }
        )
    )

    workbook = load_workbook(result.artifacts[0]["path"], data_only=False)
    worksheet = workbook["员工信息"]
    values = list(worksheet.iter_rows(min_row=1, max_row=3, max_col=5, values_only=True))
    assert values == [
        ("工号", "姓名", "部门", "基本工资(元)", "联系电话"),
        ("EMP001", "张伟", "技术部", 18000, "13800138001"),
        ("EMP002", "李娜", "产品部", 16500, "13800138002"),
    ]
    assert worksheet["B2"].font.italic is True


def test_search_executor_maps_serper_response(monkeypatch) -> None:
    monkeypatch.setattr(settings, "serper_key_id", "test-key")

    class FakeResponse:
        status_code = 200
        text = "ok"

        def json(self):
            return {
                "news": [
                    {
                        "title": "First result",
                        "link": "https://example.com/1",
                        "snippet": "Snippet 1",
                        "date": "2026-08-29",
                        "source": "Example",
                    },
                    {
                        "title": "Second result",
                        "link": "https://example.com/2",
                        "snippet": "Snippet 2",
                        "date": "2026-08-29",
                        "source": "Example",
                    },
                ]
            }

    async def fake_post(self, url, headers=None, json=None):  # type: ignore[no-untyped-def]
        assert url == "https://google.serper.dev/news"
        assert headers["X-API-KEY"] == "test-key"
        assert json["q"] == "AI chips"
        assert json["tbs"] == "qdr:d"
        return FakeResponse()

    monkeypatch.setattr(httpx.AsyncClient, "post", fake_post)

    result = asyncio.run(
        SearchExecutor().execute(
            {
                "text": "AI chips",
                "user_intent": "latest tech news",
                "parameters": {"type": "search"},
            }
        )
    )

    assert result["provider"] == "serper"
    assert result["search_type"] == "news"
    assert result["recency"] == "day"
    assert result["count"] == 2
    assert result["results"][0]["title"] == "First result"
